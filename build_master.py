# -*- coding: utf-8 -*-
"""
JBoss EAP background-validation-millis 完全技術解説書 生成マスター実行スクリプト
"""

import os
import sys
import openpyxl

from builder.sheet1_summary import build_sheet_01
from builder.sheet2_kids import build_sheet_02
from builder.sheet3_history import build_sheet_03
from builder.sheet4_arch import build_sheet_04
from builder.sheet5_compare import build_sheet_05
from builder.sheet6_ref import build_sheet_06
from builder.sheet7_troubleshoot import build_sheet_07

def main():
    print("=== JBoss EAP Excel Workbook Generation Started ===")
    
    wb = openpyxl.Workbook()
    # 初期シートを記憶
    default_sheet = wb.active
    
    print("Building Sheet 1: 01_概要・全体サマリー...")
    ws1 = build_sheet_01(wb)
    ws1.sheet_properties.tabColor = "1B365D"  # 濃紺
    
    print("Building Sheet 2: 02_小学生でもわかる解説...")
    ws2 = build_sheet_02(wb)
    ws2.sheet_properties.tabColor = "2E7D32"  # グリーン
    
    print("Building Sheet 3: 03_歴史と背景・必要性の経緯...")
    ws3 = build_sheet_03(wb)
    ws3.sheet_properties.tabColor = "E65100"  # オレンジ
    
    print("Building Sheet 4: 04_動作原理と内部機構...")
    ws4 = build_sheet_04(wb)
    ws4.sheet_properties.tabColor = "4A148C"  # パープル
    
    print("Building Sheet 5: 05_validate-on-match徹底比較...")
    ws5 = build_sheet_05(wb)
    ws5.sheet_properties.tabColor = "00695C"  # ティール
    
    print("Building Sheet 6: 06_設定リファレンスとCLI...")
    ws6 = build_sheet_06(wb)
    ws6.sheet_properties.tabColor = "1565C0"  # ブルー
    
    print("Building Sheet 7: 07_障害対策とベストプラクティス...")
    ws7 = build_sheet_07(wb)
    ws7.sheet_properties.tabColor = "C62828"  # クリムゾンレッド
    
    # デフォルトの空シートを削除
    if default_sheet in wb.worksheets:
        wb.remove(default_sheet)
        
    output_filename = "JBossEAP_background_validation_millis_完全解説書.xlsx"
    print(f"Saving workbook to {output_filename}...")
    wb.save(output_filename)
    
    print(f"=== Successfully generated {output_filename} ===")
    print(f"Total Sheets: {len(wb.sheetnames)}")
    for name in wb.sheetnames:
        print(f" - {name}")

if __name__ == "__main__":
    main()
