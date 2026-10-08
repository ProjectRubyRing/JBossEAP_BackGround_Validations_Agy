# -*- coding: utf-8 -*-
"""
共通スタイルおよびレイアウトヘルパー定義
フォント: Meiryo UI を全シート・全セルに徹底適用
"""

import math
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT_FAMILY = "Meiryo UI"

# フォント定義
FONT_SHEET_TITLE = Font(name=FONT_FAMILY, size=14, bold=True, color="FFFFFF")
FONT_SHEET_SUBTITLE = Font(name=FONT_FAMILY, size=9, italic=True, color="D9E1F2")
FONT_SEC_HEADER = Font(name=FONT_FAMILY, size=10.5, bold=True, color="FFFFFF")
FONT_SEC_SUB = Font(name=FONT_FAMILY, size=10, bold=True, color="1B365D")
FONT_TABLE_HEADER = Font(name=FONT_FAMILY, size=9.5, bold=True, color="FFFFFF")
FONT_BODY = Font(name=FONT_FAMILY, size=9, color="262626")
FONT_BODY_BOLD = Font(name=FONT_FAMILY, size=9, bold=True, color="1B365D")
FONT_BODY_MUTED = Font(name=FONT_FAMILY, size=8.5, color="595959")
FONT_CODE = Font(name=FONT_FAMILY, size=9, color="003366")
FONT_CODE_BOLD = Font(name=FONT_FAMILY, size=9, bold=True, color="003366")

# バッジ用フォント
FONT_BADGE_REC = Font(name=FONT_FAMILY, size=8.5, bold=True, color="0E6251")
FONT_BADGE_WARN = Font(name=FONT_FAMILY, size=8.5, bold=True, color="7D6608")
FONT_BADGE_DANGER = Font(name=FONT_FAMILY, size=8.5, bold=True, color="78281F")
FONT_BADGE_INFO = Font(name=FONT_FAMILY, size=8.5, bold=True, color="1B4F72")

# 塗りつぶし定義
FILL_SHEET_TITLE = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")  # 深藍
FILL_SEC_HEADER = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")    # スティールブルー
FILL_SEC_ACCENT = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")    # 淡いブルー
FILL_TABLE_HEADER = PatternFill(start_color="203764", end_color="203764", fill_type="solid")  # 濃紺ヘッダー
FILL_TH_CHARCOAL = PatternFill(start_color="333F48", end_color="333F48", fill_type="solid")  # チャコールグレー
FILL_ZEBRA_EVEN = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")   # 偶数行（極淡ブルー）
FILL_ZEBRA_ODD = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")    # 奇数行（白）
FILL_CODE_BLOCK = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")   # コード背景

# バッジ背景
FILL_BADGE_REC = PatternFill(start_color="E8F8F5", end_color="E8F8F5", fill_type="solid")     # ミントグリーン
FILL_BADGE_WARN = PatternFill(start_color="FEF9E7", end_color="FEF9E7", fill_type="solid")    # ライトイエロー
FILL_BADGE_DANGER = PatternFill(start_color="FDEDEC", end_color="FDEDEC", fill_type="solid")  # ライトレッド
FILL_BADGE_INFO = PatternFill(start_color="EBF5FB", end_color="EBF5FB", fill_type="solid")    # ライトブルー

# 罫線定義
BORDER_THIN = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9")
)
BORDER_HEADER = Border(
    left=Side(style="thin", color="B0C4DE"),
    right=Side(style="thin", color="B0C4DE"),
    top=Side(style="medium", color="1B365D"),
    bottom=Side(style="medium", color="1B365D")
)
BORDER_TOTAL = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="1B365D"),
    bottom=Side(style="double", color="1B365D")
)

# 配置定義
ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
ALIGN_LEFT_TOP = Alignment(horizontal="left", vertical="top", wrap_text=True)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_CENTER_TOP = Alignment(horizontal="center", vertical="top", wrap_text=True)
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center", wrap_text=True)
ALIGN_RIGHT_TOP = Alignment(horizontal="right", vertical="top", wrap_text=True)
ALIGN_CODE = Alignment(horizontal="left", vertical="top", wrap_text=True)


def setup_sheet_view(ws):
    """グリッド線の表示を保証"""
    ws.views.sheetView[0].showGridLines = True


def set_col_widths(ws, col_width_dict):
    """列幅を設定"""
    for col_letter, width in col_width_dict.items():
        ws.column_dimensions[col_letter].width = width


def auto_fit_row_height(ws, row_idx, col_widths_map, min_height=20):
    """
    行内のテキスト量と改行数から適切な行の高さを算出し設定
    col_widths_map: {col_letter: width, ...}
    """
    max_lines = 1
    for col_letter, width in col_widths_map.items():
        c_idx = openpyxl.utils.column_index_from_string(col_letter)
        cell_val = ws.cell(row=row_idx, column=c_idx).value
        if cell_val is not None:
            text = str(cell_val)
            lines = text.split("\n")
            line_count = 0
            for line in lines:
                char_len = sum(2 if ord(ch) > 127 else 1 for ch in line)
                effective_width = max(10, width * 1.7)
                sub_lines = max(1, math.ceil(char_len / effective_width))
                line_count += sub_lines
            if line_count > max_lines:
                max_lines = line_count
    calculated_height = max(min_height, max_lines * 14.5 + 7)
    ws.row_dimensions[row_idx].height = calculated_height


def render_title_banner(ws, title_text, subtitle_text, max_col_letter="E"):
    """シート先頭のタイトルバナーを描画（結合セル全域にスタイルを適用）"""
    start_c_idx = openpyxl.utils.column_index_from_string("B")
    end_c_idx = openpyxl.utils.column_index_from_string(max_col_letter)
    
    for c_idx in range(start_c_idx, end_c_idx + 1):
        c2 = ws.cell(row=2, column=c_idx)
        c2.fill = FILL_SHEET_TITLE
        c3 = ws.cell(row=3, column=c_idx)
        c3.fill = FILL_SHEET_TITLE

    ws.merge_cells(f"B2:{max_col_letter}2")
    ws.merge_cells(f"B3:{max_col_letter}3")
    
    t_cell = ws["B2"]
    t_cell.value = title_text
    t_cell.font = FONT_SHEET_TITLE
    t_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    
    s_cell = ws["B3"]
    s_cell.value = subtitle_text
    s_cell.font = FONT_SHEET_SUBTITLE
    s_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    
    ws.row_dimensions[2].height = 28
    ws.row_dimensions[3].height = 18


def render_section_header(ws, row_idx, section_text, max_col_letter="E"):
    """セクション見出しを描画（結合セル全域にスタイルを適用）"""
    start_c_idx = openpyxl.utils.column_index_from_string("B")
    end_c_idx = openpyxl.utils.column_index_from_string(max_col_letter)
    
    for c_idx in range(start_c_idx, end_c_idx + 1):
        cell = ws.cell(row=row_idx, column=c_idx)
        cell.fill = FILL_SEC_HEADER

    ws.merge_cells(f"B{row_idx}:{max_col_letter}{row_idx}")
    cell = ws[f"B{row_idx}"]
    cell.value = f" ■ {section_text}"
    cell.font = FONT_SEC_HEADER
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row_idx].height = 24


def render_table_headers(ws, row_idx, headers, start_col="B"):
    """テーブルヘッダー行を描画"""
    start_c_idx = openpyxl.utils.column_index_from_string(start_col)
    for i, header_text in enumerate(headers):
        col_idx = start_c_idx + i
        cell = ws.cell(row=row_idx, column=col_idx)
        cell.value = header_text
        cell.font = FONT_TABLE_HEADER
        cell.fill = FILL_TABLE_HEADER
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER_HEADER
    ws.row_dimensions[row_idx].height = 24
